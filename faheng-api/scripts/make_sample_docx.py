"""生成测试用"乙方版本"合同（故意埋坑，条款取自 ERP 审核原型的 before 字段）

运行：python scripts/make_sample_docx.py
输出：data/sample_contract.docx
"""
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt

# (条款标题, 条款正文)——每条都命中规则库
CLAUSES = [
    ("第一条 定义与合同文件",
     "1.3 前述文件不一致时，以更新时间较晚的乙方产品规则或服务政策为准。"),
    ("第二条 服务内容与项目范围",
     "2.1 许可仅限甲方本部使用，甲方关联公司、分支机构及外部合作方使用应另行购买许可。\n"
     "2.3 乙方判断超出标准范围的工作，按实际人天另行收费。\n"
     "2.4 乙方可根据产品规划对模块和功能进行升级、合并、替换或停止维护。"),
    ("第三条 私有云部署条件与责任边界",
     "3.2 凡无法明确证明由软件代码直接引起的故障，均视为甲方环境原因。\n"
     "3.3 乙方可在必要范围内调取系统日志、配置文件、业务数据样本和数据库信息，无需逐次取得甲方确认。"),
    ("第四条 项目计划、交付与验收",
     "4.1 预计实施周期为六个月。该周期为预估时间，不构成乙方的确定交付承诺。\n"
     "4.3 乙方通知甲方验收后，甲方应在5个工作日内一次性书面提出全部异议；逾期未回复、甲方开始试运行或任一用户登录软件的，均视为项目全部验收合格。"),
    ("第五条 合同金额与支付",
     "5.1 甲方应在合同签署后5个工作日内支付首年订阅费、全部实施费及全部数据迁移费。\n"
     "5.2 所有费用一经支付不予退还。\n"
     "5.3 乙方有权根据人工、基础软件或监管成本变化调整后续年度订阅价格，并提前15日通知甲方；甲方继续使用即视为接受调整。\n"
     "5.4 每逾期一日按应付未付金额的0.1%支付违约金。逾期超过3日，乙方可暂停服务。"),
    ("第六条 订阅许可与账号管理",
     "6.2 超出许可的，甲方应按乙方届时公开价格补缴费用并加付50%的许可管理费。"),
    ("第七条 运维支持与服务等级",
     "7.1 重大故障的首次响应时间不超过4小时。首次响应仅指乙方确认收到问题。\n"
     "7.2 因计划维护、版本升级、安全攻击或乙方合理控制范围之外的原因造成的中断，不计入不可用时间。\n"
     "7.3 服务期限抵扣是甲方就服务不可用可获得的唯一救济。"),
    ("第八条 数据处理、迁移与备份",
     "8.2 乙方可将甲方数据去标识化或汇总后用于算法训练、产品改进、行业分析、商业展示及其他经营目的。\n"
     "8.3 因甲方数据引发的投诉、调查或索赔由甲方承担。\n"
     "8.5 乙方可根据排障需要在其或合作伙伴环境中临时保存数据副本。\n"
     "8.6 甲方应在终止后7日内完成下载，逾期乙方可删除数据且不承担责任。"),
    ("第九条 网络与信息安全",
     "9.2 乙方将在完成内部核实并确认事件对甲方产生实质影响后通知甲方。\n"
     "9.3 因甲方环境或双方原因混合导致的安全事件，乙方不承担责任。\n"
     "9.4 甲方开展安全测试应至少提前30日取得乙方书面同意。"),
    ("第十条 知识产权",
     "10.2 无论是否由甲方付费提出需求，定制成果知识产权均归乙方所有。\n"
     "10.4 前述修改、替换或退款措施为乙方承担的全部责任。"),
    ("第十一条 保密",
     "11.2 乙方仅对甲方明确标注为\"保密\"的书面资料承担保密义务。"),
    ("第十二条 分包与第三方服务",
     "12.1 乙方可将实施、运维、客服、数据处理委托第三方，无需甲方另行同意。"),
    ("第十四条 暂停服务",
     "14.1 乙方认为可能影响软件安全时，可立即暂停服务且无需承担责任。"),
    ("第十五条 陈述保证与责任限制",
     "15.3 乙方在本合同项下的累计责任总额不超过导致索赔事件发生前连续三个月甲方已支付的软件订阅费。\n"
     "15.4 因甲方依据软件报表、算法或计算结果作出决策产生的损失，由甲方自行承担。"),
    ("第十六条 合同期限、续订与终止",
     "16.1 服务期自乙方交付管理员账号之日起计算，实施期和试运行期均计入服务期。\n"
     "16.2 服务期届满前60日甲方未提出不续订，本合同自动续订一年。\n"
     "16.3 甲方确需提前终止的，应提前90日书面通知乙方，并支付剩余合同期全部订阅费及未结费用。"),
    ("第十八条 不可抗力",
     "18.1 不可抗力包括网络攻击、病毒、云服务商或基础软件供应商故障。"),
    ("第十九条 通知与电子证据",
     "19.1 乙方向管理后台发送信息后即视为送达。"),
    ("第二十条 法律适用与争议解决",
     "20.1 协商不成的，任何一方均应向乙方住所地有管辖权的人民法院提起诉讼。"),
    ("第二十一条 其他",
     "21.1 乙方可将本合同转让给其关联公司或业务承继方，并通知甲方。"),
]

PREAMBLE = [
    "ERP软件订阅及私有云部署服务合同",
    "合同编号：YH-ERP-2026-0829",
    "",
    "甲方（采购方）：北京示例制造有限公司",
    "住所地：北京市朝阳区示例路1号",
    "",
    "乙方（供应方）：上海示例软件技术有限公司",
    "住所地：上海市浦东新区示例大道88号",
    "",
    "鉴于甲方拟采购乙方ERP软件订阅及私有云部署服务，双方经友好协商达成如下协议：",
]

TAIL = [
    "（以下无正文，为签署页）",
    "",
    "甲方（盖章）：北京示例制造有限公司",
    "授权代表（签字）：",
    "签署日期：2026年　月　日",
    "",
    "乙方（盖章）：上海示例软件技术有限公司",
    "授权代表（签字）：",
    "签署日期：2026年　月　日",
]


def main() -> None:
    import sys

    sys.stdout.reconfigure(encoding="utf-8")
    # 让脚本可以直接 python scripts/make_sample_docx.py 运行
    sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
    doc = Document()

    # 标题
    p = doc.add_paragraph()
    run = p.add_run(PREAMBLE[0])
    run.bold = True
    run.font.size = Pt(16)
    _set_font(run, "黑体")
    # 标题占 preamble 第一行，其余行从第二行起
    for line in PREAMBLE[1:]:
        if line.strip():
            p = doc.add_paragraph()
            run = p.add_run(line)
            _set_font(run)

    for title, body in CLAUSES:
        p = doc.add_paragraph()
        run = p.add_run(title)
        run.bold = True
        _set_font(run)
        for line in body.splitlines():
            p = doc.add_paragraph()
            run = p.add_run(line)
            _set_font(run)

    for line in TAIL:
        if line.strip():
            p = doc.add_paragraph()
            run = p.add_run(line)
            _set_font(run)

    out = Path(__file__).resolve().parent.parent / "data" / "sample_contract.docx"
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out))
    print(f"生成测试合同: {out}")

    # 自检切分
    from app.parser.clause_splitter import split_clauses
    from app.parser.docx_parser import extract_docx_paragraphs

    lines = extract_docx_paragraphs(out)
    clauses, pre, tail = split_clauses(lines)
    print(f"切分结果: {len(clauses)} 条条款, preamble {len(pre)} 行, tail {len(tail)} 行")
    for c in clauses:
        print(f"  {c.clause_id} {c.clause_no} {c.title} [{c.start_idx}-{c.end_idx}]")


def _set_font(run, east_asia: str = "宋体"):
    run.font.name = "Times New Roman"
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = rpr.makeelement(qn("w:rFonts"), {})
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), east_asia)


if __name__ == "__main__":
    main()
