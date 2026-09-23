"""DOCX 解析：python-docx 段落提取（表格文本跳过，导出时原样保留）

返回 `ParagraphInfo` 列表：
- `text`：原段落纯文本（与原解析一致，供切分/导出锚点使用）
- `html`：包含段落样式与 Run 级样式的安全 HTML，供前端展示使用

下标与原解析保持一一对应（空段落也保留为 `<p><br></p>`），
因此 `start_idx / end_idx` 锚点不受影响。
"""
from __future__ import annotations

import html
import re
from typing import NamedTuple

from docx import Document
from docx.document import Document as _DocxDocument
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph


class ParagraphInfo(NamedTuple):
    text: str
    html: str


# python-docx 的 WD_ALIGN_PARAGRAPH 整型映射
_ALIGN_MAP = {
    0: "left",     # WD_ALIGN_PARAGRAPH.LEFT
    1: "center",   # WD_ALIGN_PARAGRAPH.CENTER
    2: "right",    # WD_ALIGN_PARAGRAPH.RIGHT
    3: "justify",  # WD_ALIGN_PARAGRAPH.JUSTIFY
}


def _em_to_pt(value) -> float | None:
    """python-docx 的 Pt/Emu 等单位对象 → pt 数值；非数字返回 None。"""
    try:
        return float(value.pt)
    except Exception:
        return None


def _color_hex(color) -> str | None:
    """run.font.color.rgb 是 RGB 对象；部分场景是 theme 类型，没法可靠取色，返回 None。"""
    try:
        rgb = color.rgb
        if rgb is None:
            return None
        return f"#{str(rgb)}"
    except Exception:
        return None


def _run_to_html(run) -> str:
    """把单个 run 渲染成 HTML 片段，含粗体/斜体/下划线/颜色/字号/字体。"""
    raw = run.text or ""
    if not raw:
        return ""
    text = html.escape(raw).replace("\n", "<br>")
    if not text.strip():
        return ""

    styles: list[str] = []
    font = run.font
    if font.size is not None:
        pt = _em_to_pt(font.size)
        if pt:
            styles.append(f"font-size:{pt:g}pt")
    if font.name:
        styles.append(f"font-family:{html.escape(font.name)},serif")
    color = _color_hex(font.color)
    if color:
        styles.append(f"color:{color}")
    if font.highlight_color is not None:
        # 高亮色（python-docx 返回 WD_COLOR_INDEX 枚举）
        styles.append("background-color:#fff59d")

    style_attr = f' style="{";".join(styles)}"' if styles else ""
    inner = f'<span{style_attr}>{text}</span>' if style_attr else text

    # 粗/斜/下划线三种简单装饰
    if font.bold:
        inner = f"<strong>{inner}</strong>"
    if font.italic:
        inner = f"<em>{inner}</em>"
    if font.underline:
        inner = f"<u>{inner}</u>"
    return inner


def _paragraph_to_html(p: Paragraph) -> str:
    """段落 → `<p style="text-align:...">…runs…</p>`，空段为 `<p><br></p>`。"""
    runs_html = "".join(_run_to_html(r) for r in p.runs)
    if not runs_html:
        # 纯空段或纯空白段：占位 <br>，保留锚点
        body = "<br>"
    else:
        body = runs_html

    align = None
    try:
        align_value = p.paragraph_format.alignment
        if align_value is not None:
            align = _ALIGN_MAP.get(int(align_value))
    except Exception:
        align = None

    style_attr = f' style="text-align:{align}"' if align else ""
    return f"<p{style_attr}>{body}</p>"


def extract_docx_paragraphs(src) -> list[ParagraphInfo]:
    """返回 body 级段落（含空段），下标即导出回填锚点。

    `src` 可为：文件路径（str / Path）、或 file-like 对象（如 ``io.BytesIO``）。
    python-docx 的 ``Document(src)`` 原生接受这两类入参。
    """
    doc: _DocxDocument = Document(src)
    result: list[ParagraphInfo] = []
    for p in doc.paragraphs:
        text = p.text
        html_str = _paragraph_to_html(p)
        result.append(ParagraphInfo(text=text, html=html_str))
    return result
